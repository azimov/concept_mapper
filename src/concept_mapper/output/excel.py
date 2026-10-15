"""Excel workbook output with highlight of missed and added concepts."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from concept_mapper.models import MappingResult, ValidationReport

_RED = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
_RED_FONT = Font(color="FF9C0006")
_YELLOW = PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid")
_YELLOW_FONT = Font(color="FF9C6500")
_GREEN = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
_GREEN_FONT = Font(color="FF006100")
_ORANGE = PatternFill(start_color="FFF8CBAD", end_color="FFF8CBAD", fill_type="solid")
_ORANGE_FONT = Font(color="FF833C0B")
_IMPACT_STYLE = {
    "trivial": (_GREEN, _GREEN_FONT),
    "needs review": (_YELLOW, _YELLOW_FONT),
    "not removable alone": (_ORANGE, _ORANGE_FONT),
}
_HEADER_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def _autosize(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_header(ws, columns: list[str], widths: list[int]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    _autosize(ws, widths)


def _std_refs(ids: list[int], report: ValidationReport) -> tuple[str, str, str]:
    names: list[str] = []
    codes: list[str] = []
    for cid in ids:
        standard = report.result.standard_concepts.get(cid)
        if standard is not None:
            names.append(standard.concept.concept_name)
            codes.append(standard.concept.concept_code)
    return (
        "; ".join(str(i) for i in ids),
        "; ".join(codes),
        "; ".join(names),
    )


def _write_excluded_sheets(wb: Workbook, report: ValidationReport, excluded: MappingResult) -> None:
    included_ids = set(report.result.standard_ids)
    ws = wb.create_sheet("Excluded Source Codes")
    _write_header(
        ws,
        [
            "input_code",
            "matched_code",
            "match_type",
            "source_concept_name",
            "source_vocabulary_id",
            "standard_concept_ids",
            "standard_concept_codes",
            "standard_concept_names",
            "status",
            "notes",
        ],
        [14, 14, 12, 42, 20, 18, 24, 48, 12, 50],
    )
    for match in excluded.source_matches:
        src = match.source_concept
        ids = match.standard_concept_ids
        found = [
            excluded.standard_concepts[i].concept
            for i in ids
            if i in excluded.standard_concepts
        ]
        codes = [c.concept_code for c in found]
        names = [c.concept_name for c in found]
        ws.append(
            [
                match.input_code,
                match.matched_code or "",
                match.match_type,
                src.concept_name if src else "",
                src.vocabulary_id if src else "",
                "; ".join(str(i) for i in ids),
                "; ".join(codes),
                "; ".join(names),
                match.status,
                "; ".join(match.notes),
            ]
        )
        if match.status != "mapped":
            for cell in ws[ws.max_row]:
                cell.fill = _RED
                cell.font = _RED_FONT

    ws2 = wb.create_sheet("Excluded Concepts")
    _write_header(
        ws2,
        [
            "concept_id",
            "concept_code",
            "concept_name",
            "vocabulary_id",
            "domain_id",
            "excluded_source_codes",
            "also_in_included_set",
        ],
        [18, 14, 48, 14, 14, 40, 20],
    )
    for cid in excluded.standard_ids:
        standard = excluded.standard_concepts[cid]
        c = standard.concept
        overlap = cid in included_ids
        ws2.append(
            [
                c.concept_id,
                c.concept_code,
                c.concept_name,
                c.vocabulary_id,
                c.domain_id,
                "; ".join(standard.source_codes),
                "yes" if overlap else "no",
            ]
        )
        if overlap:
            for cell in ws2[ws2.max_row]:
                cell.fill = _YELLOW
                cell.font = _YELLOW_FONT


def _write_standard_vs_source(wb: Workbook, report: ValidationReport) -> None:
    ws = wb.create_sheet("Standard vs Source")
    _write_header(
        ws,
        [
            "standard_concept_id",
            "standard_concept_name",
            "standard_record_count",
            "source_concept_id",
            "source_concept_code",
            "source_concept_name",
            "source_vocabulary_id",
            "source_record_count",
            "in_input",
        ],
        [18, 48, 18, 18, 16, 48, 18, 18, 10],
    )
    for cid in report.result.standard_ids:
        standard = report.result.standard_concepts[cid]
        for src in standard.source_counts:
            ws.append(
                [
                    cid,
                    standard.concept.concept_name,
                    standard.record_count if standard.record_count is not None else "",
                    src.concept_id,
                    src.concept_code,
                    src.concept_name,
                    src.vocabulary_id,
                    src.record_count if src.record_count is not None else "",
                    "yes" if src.in_input else "no",
                ]
            )
            if not src.in_input:
                for cell in ws[ws.max_row]:
                    cell.fill = _YELLOW
                    cell.font = _YELLOW_FONT


def _write_overbroad_sources(wb: Workbook, report: ValidationReport) -> None:
    ws = wb.create_sheet("Overbroad Source Codes")
    _write_header(
        ws,
        [
            "source_concept_id",
            "source_concept_code",
            "source_concept_name",
            "source_vocabulary_id",
            "source_record_count",
            "via_standard_concept_id",
            "via_standard_concept_name",
            "standard_record_count",
            "introducing_input_codes",
            "input_alternative_standards",
            "input_one_to_many",
            "also_mapped_via_in_set_standards",
            "removal_impact",
        ],
        [18, 16, 48, 18, 18, 22, 48, 20, 24, 40, 16, 32, 22],
    )
    standards = report.result.standard_concepts

    def source_count(o) -> int | None:
        for src in standards[o.standard_concept_id].source_counts:
            if src.concept_id == o.source_concept_id:
                return src.record_count
        return None

    rows = [(o, source_count(o)) for o in report.overbroad_sources]
    rows.sort(key=lambda r: (-(r[1] or 0), r[0].source_concept_code, r[0].standard_concept_id))
    for o, count in rows:
        standard = standards[o.standard_concept_id]
        alternatives = "; ".join(
            f"{code} -> {', '.join(str(i) for i in alts) if alts else 'none (sole mapping)'}"
            for code, alts in o.alternative_standards.items()
        )
        ws.append(
            [
                o.source_concept_id,
                o.source_concept_code,
                o.source_concept_name,
                o.source_vocabulary_id,
                count if count is not None else "",
                o.standard_concept_id,
                standard.concept.concept_name,
                standard.record_count if standard.record_count is not None else "",
                "; ".join(o.introducing_input_codes),
                alternatives,
                "yes" if o.input_one_to_many else "no",
                "; ".join(str(i) for i in o.other_in_set_standard_ids),
                o.removal_impact,
            ]
        )
        fill, font = _IMPACT_STYLE[o.removal_impact]
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.font = font


def write_workbook(
    path: str | Path,
    report: ValidationReport,
    excluded_result: MappingResult | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ── Source Codes sheet ────────────────────────────────────────────
    ws = wb.active
    ws.title = "Source Codes"
    _write_header(
        ws,
        [
            "input_code",
            "matched_code",
            "match_type",
            "source_concept_id",
            "source_concept_name",
            "source_vocabulary_id",
            "source_domain_id",
            "standard_concept_ids",
            "standard_concept_codes",
            "standard_concept_names",
            "value_concept_ids",
            "source_record_count",
            "status",
            "notes",
        ],
        [14, 14, 12, 18, 42, 20, 16, 18, 24, 48, 18, 18, 12, 50],
    )
    for match in report.result.source_matches:
        src = match.source_concept
        std_ids, std_codes, std_names = _std_refs(match.standard_concept_ids, report)
        row = [
            match.input_code,
            match.matched_code or "",
            match.match_type,
            src.concept_id if src else "",
            src.concept_name if src else "",
            src.vocabulary_id if src else "",
            src.domain_id if src else "",
            std_ids,
            std_codes,
            std_names,
            "; ".join(str(v) for v in match.value_concept_ids),
            match.source_record_count if match.source_record_count is not None else "",
            match.status,
            "; ".join(match.notes),
        ]
        ws.append(row)
        if match.status != "mapped":
            for cell in ws[ws.max_row]:
                cell.fill = _RED
                cell.font = _RED_FONT

    # ── Standard Concepts sheet ───────────────────────────────────────
    ws2 = wb.create_sheet("Standard Concepts")
    _write_header(
        ws2,
        [
            "concept_id",
            "concept_code",
            "concept_name",
            "vocabulary_id",
            "domain_id",
            "concept_class_id",
            "source_codes",
            "extra_source_codes",
            "record_count",
            "source_concept_count",
            "is_overbroad",
            "recommendation",
        ],
        [18, 14, 48, 14, 14, 18, 40, 40, 14, 20, 14, 24],
    )
    recommendation_by_id: dict[int, str] = {}
    for standard in report.exclusions:
        cid = standard.concept.concept_id
        surviving = report.surviving_extra_codes(cid)
        if surviving:
            detail = ", ".join(
                f"{code} still via {', '.join(str(i) for i in ids)}"
                for code, ids in sorted(surviving.items())
            )
            recommendation_by_id[cid] = f"exclude (partial: {detail})"
        else:
            recommendation_by_id[cid] = "exclude"
    for standard in report.needs_review:
        recommendation_by_id[standard.concept.concept_id] = "needs review"

    for cid in report.result.standard_ids:
        standard = report.result.standard_concepts[cid]
        c = standard.concept
        ws2.append(
            [
                c.concept_id,
                c.concept_code,
                c.concept_name,
                c.vocabulary_id,
                c.domain_id,
                c.concept_class_id,
                "; ".join(standard.source_codes),
                "; ".join(standard.extra_source_codes),
                standard.record_count if standard.record_count is not None else "",
                len(standard.source_counts) if standard.source_counts else "",
                "yes" if standard.is_overbroad else "no",
                recommendation_by_id.get(cid, ""),
            ]
        )
        if standard.is_overbroad:
            for cell in ws2[ws2.max_row]:
                cell.fill = _YELLOW
                cell.font = _YELLOW_FONT

    # ── Overbroad Source Codes sheet ──────────────────────────────────
    _write_overbroad_sources(wb, report)

    # ── Standard vs Source sheet ──────────────────────────────────────
    if any(s.source_counts for s in report.result.standard_concepts.values()):
        _write_standard_vs_source(wb, report)

    # ── Summary sheet ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")
    _write_header(ws3, ["section", "detail"], [24, 100])
    total = len(report.result.source_matches)
    mapped = total - len(report.missed)
    ws3.append(["total input codes", str(total)])
    ws3.append(["mapped", str(mapped)])
    ws3.append(["missed", str(len(report.missed))])
    ws3.append(["standard concepts in set", str(len(report.result.standard_ids))])
    ws3.append(["added / overbroad", str(len(report.added))])
    ws3.append(["suggested exclusions", str(len(report.exclusions))])
    ws3.append(["needs review", str(len(report.needs_review))])
    overbroad_codes = {o.source_concept_id for o in report.overbroad_sources}
    ws3.append(["overbroad source codes", str(len(overbroad_codes))])
    for impact in _IMPACT_STYLE:
        n = sum(1 for o in report.overbroad_sources if o.removal_impact == impact)
        ws3.append([f"  removal impact: {impact}", str(n)])
    ws3.append([])
    ws3.append(["Missed source codes", ""])
    for m in report.missed:
        reason = "not found" if m.match_type == "not_found" else "no standard mapping"
        ws3.append([m.input_code, reason])
    ws3.append([])
    ws3.append(["Added / overbroad standard concepts", ""])
    for standard in report.added:
        c = standard.concept
        ws3.append(
            [f"{c.concept_id} ({c.concept_code})",
             f"{c.concept_name} -> extra codes: {', '.join(standard.extra_source_codes)}"]
        )
    ws3.append([])
    ws3.append(["Suggested exclusions", ""])
    for standard in report.exclusions:
        c = standard.concept
        ws3.append([f"{c.concept_id} ({c.concept_code})", c.concept_name])
    ws3.append([])
    ws3.append(["Needs review (overbroad but required)", ""])
    for standard in report.needs_review:
        c = standard.concept
        suggestions = report.descendant_suggestions.get(c.concept_id, [])
        desc = "; ".join(f"{d.concept_id} {d.concept_name}" for d in suggestions)
        ws3.append([f"{c.concept_id} ({c.concept_code})", f"{c.concept_name}; consider: {desc}"])

    if excluded_result is not None:
        _write_excluded_sheets(wb, report, excluded_result)

    wb.save(path)
    return path
