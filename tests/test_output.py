from openpyxl import load_workbook

from concept_mapper.mapping import map_codes
from concept_mapper.output.concept_set import build_concept_set, write_concept_set
from concept_mapper.output.excel import write_workbook
from concept_mapper.validation import validate_result

SOURCE_VOCABULARIES = ["ICD10CM"]


def _report(repo):
    result = map_codes(repo, ["E11.9", "C01.9", "I21.0", "Z01.0"], SOURCE_VOCABULARIES, "SNOMED")
    return validate_result(repo, result, SOURCE_VOCABULARIES, "SNOMED")


def test_concept_set_json_includes_exclusions(repo):
    report = _report(repo)
    data = build_concept_set(report.result, report.exclusions)
    assert len(data["items"]) == len(report.result.standard_ids) + len(report.exclusions)
    concepts = {i["concept"]["CONCEPT_ID"]: i for i in data["items"]}
    assert concepts[2004]["isExcluded"] is True
    assert concepts[2001]["isExcluded"] is False
    assert concepts[2001]["concept"]["CONCEPT_CODE"] == "44054006"


def test_write_concept_set(repo, tmp_path):
    report = _report(repo)
    path = write_concept_set(
        tmp_path / "concept_set.json", report.result, report.exclusions, name="x"
    )
    assert path.exists()
    assert '"items"' in path.read_text()


def test_write_workbook(repo, tmp_path):
    report = _report(repo)
    path = write_workbook(tmp_path / "concept_set.xlsx", report)
    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Source Codes", "Standard Concepts", "Overbroad Source Codes", "Summary"
    ]


def _sheet_rows(ws):
    header = [c.value for c in ws[1]]
    return [
        dict(zip(header, (c.value for c in row), strict=False))
        for row in ws.iter_rows(min_row=2)
    ]


def test_overbroad_source_codes_sheet(repo, tmp_path):
    report = _report(repo)
    wb = load_workbook(write_workbook(tmp_path / "concept_set.xlsx", report))
    rows = {r["source_concept_code"]: r for r in _sheet_rows(wb["Overbroad Source Codes"])}
    assert set(rows) == {"E11", "I21"}
    assert rows["I21"]["via_standard_concept_id"] == 2004
    assert rows["I21"]["introducing_input_codes"] == "I21.0"
    assert rows["I21"]["input_alternative_standards"] == "I21.0 -> 2003"
    assert rows["I21"]["input_one_to_many"] == "yes"
    assert rows["I21"]["removal_impact"] == "trivial"
    assert rows["E11"]["removal_impact"] == "needs review"
    assert rows["E11"]["source_record_count"] is None


def test_overbroad_source_codes_with_counts(fixtures_dir, tmp_path):
    from concept_mapper.frequency import count_frequencies
    from concept_mapper.repository import DuckDBRepository

    repo = DuckDBRepository(csv_dir=str(fixtures_dir))
    try:
        repo.load_table("condition_occurrence", str(fixtures_dir / "CONDITION_OCCURRENCE.csv"))
        result = map_codes(repo, ["E11.9", "K50.0", "K50.1"], SOURCE_VOCABULARIES, "SNOMED")
        report = validate_result(repo, result, SOURCE_VOCABULARIES, "SNOMED")
        count_frequencies(repo, result, source_vocabularies=SOURCE_VOCABULARIES)
        wb = load_workbook(write_workbook(tmp_path / "concept_set.xlsx", report))
    finally:
        repo.close()

    rows = _sheet_rows(wb["Overbroad Source Codes"])
    by_key = {(r["source_concept_code"], r["via_standard_concept_id"]): r for r in rows}
    assert by_key[("E11", 2001)]["source_record_count"] == 1
    assert by_key[("K50", 2006)]["source_record_count"] == 2
    assert by_key[("K50", 2006)]["also_mapped_via_in_set_standards"] == "2007"
    assert by_key[("K50", 2006)]["removal_impact"] == "not removable alone"
    assert rows[0]["source_record_count"] >= rows[-1]["source_record_count"]
